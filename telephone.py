# ============================================================
# FORCE UTF-8 OUTPUT (fixes UnicodeEncodeError when this script is
# launched via subprocess, e.g. from app5.py). Must run before any print().
# ============================================================
import sys
import io

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    sys.stderr.reconfigure(encoding="utf-8", errors="replace")
else:
    sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8", errors="replace")
    sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding="utf-8", errors="replace")

import os
import re
import time
import random
import logging
import platform
import atexit

import pandas as pd
from openpyxl import load_workbook
from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import TimeoutException

# Get the file path from the command-line argument
if len(sys.argv) < 2:
    print("Error: No file path provided!")
    print("Usage: python3 telephone.py <excel_file_path>")
    sys.exit(1)

file_path = sys.argv[1]
# Only private persons are processed. Companies are skipped on purpose.
sheets = ["Private "]

# The log file is written next to the input Excel file (inside this run's
# own session folder) so several people running at once don't overwrite
# each other's logs.
log_path = os.path.join(os.path.dirname(file_path) or '.', 'telephone_debug.log')
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s',
    handlers=[logging.FileHandler(log_path), logging.StreamHandler()]
)
logger = logging.getLogger(__name__)

print(f"Received file path: {file_path}")
print(f"File exists: {os.path.exists(file_path)}")
if os.path.exists(file_path):
    print(f"File size: {os.path.getsize(file_path)} bytes")
    print(f"File extension: {os.path.splitext(file_path)[1]}")
else:
    print("ERROR: File does not exist!")
    sys.exit(1)

# ============================================================
# VIRTUAL DISPLAY (Xvfb) FOR SERVER DEPLOYMENT
# ============================================================
# No Chrome --headless on the server (headless is flagged as a bot more
# often). Instead a virtual display is started and Chrome runs normally.
# Requires: `apt install xvfb` and `pip install pyvirtualdisplay`.
# On Windows this block is skipped, so a real Chrome window opens.
virtual_display = None
used_headless_fallback = False

if platform.system() == "Linux":
    try:
        from pyvirtualdisplay import Display
        virtual_display = Display(visible=False, size=(1920, 1080))
        virtual_display.start()
        atexit.register(virtual_display.stop)
        print("🖥️ Virtual display (Xvfb) started for server environment")
    except Exception as e:
        print(f"⚠️ Could not start virtual display ({e}) - falling back to --headless=new")
        used_headless_fallback = True

options = webdriver.ChromeOptions()
options.add_argument("--no-sandbox")
options.add_argument("--disable-dev-shm-usage")
options.add_argument("--disable-gpu")
options.add_argument("--disable-blink-features=AutomationControlled")
options.add_argument("--disable-extensions")
options.add_experimental_option("excludeSwitches", ["enable-automation", "enable-logging"])
options.add_experimental_option('useAutomationExtension', False)

if used_headless_fallback:
    options.add_argument("--headless=new")

user_agents = [
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
]
options.add_argument(f"--user-agent={random.choice(user_agents)}")

driver = None


def _cleanup():
    """Makes sure Chrome never stays hanging in memory, however the
    script exits."""
    try:
        if driver is not None:
            driver.quit()
    except Exception:
        pass


atexit.register(_cleanup)

try:
    # Service() with no arguments: Selenium picks and caches the matching
    # ChromeDriver itself (no webdriver-manager lock problems).
    driver = webdriver.Chrome(service=Service(), options=options)
    driver.execute_script("Object.defineProperty(navigator, 'webdriver', {get: () => undefined})")
    print("🔧 Chrome driver setup completed")
except Exception as e:
    print(f"❌ Failed to start Chrome driver: {e}")
    sys.exit(1)

print("\n🌐 ACCESSING 1881.NO...")
print("=" * 60)

try:
    driver.get("https://www.1881.no/")
    print("✅ Page request sent to 1881.no")
    time.sleep(random.uniform(3, 6))
except Exception as e:
    print(f"❌ Error accessing 1881.no: {e}")
    driver.quit()
    sys.exit(1)


def handle_consent_popup():
    """Handle the consent popup that appears on 1881.no"""
    print("\n🔍 HANDLING CONSENT POPUP...")
    print("=" * 50)
    try:
        time.sleep(3)
        if "give us your consent" in driver.page_source.lower():
            print("🔍 Consent popup detected, attempting to close...")
            button_selectors = [
                "button[class*='agree']",
                ".consent-button",
                "[data-consent='accept']",
            ]
            for selector in button_selectors:
                try:
                    button = WebDriverWait(driver, 5).until(
                        EC.element_to_be_clickable((By.CSS_SELECTOR, selector))
                    )
                    driver.execute_script("arguments[0].click();", button)
                    print("✅ Consent popup closed successfully!")
                    time.sleep(2)
                    return True
                except Exception:
                    continue
            print("⚠️ Could not automatically close consent popup")
            return False
        print("✅ No consent popup detected")
        return True
    except Exception as e:
        print(f"❌ Error handling consent popup: {e}")
        return False


def verify_person_match(page_name, search_name, zip_code):
    """Verify if the person on the page matches the search criteria"""
    try:
        page_n = page_name.lower().strip()
        search_n = search_name.lower().strip()
        name_match = (search_n in page_n or page_n in search_n
                      or search_n.split()[0] in page_n)
        if not name_match:
            print(f"⚠️ Name mismatch: expected '{search_name}', found '{page_name}'")
            return False
        print(f"✅ Name match confirmed: '{page_name}'")
        return True
    except Exception as e:
        print(f"⚠️ Could not verify person match: {e}")
        return True


def extract_postal_code_from_address(address_text):
    """Extract postal code from address text"""
    try:
        m = re.search(r'(\d{4})', address_text)
        if m:
            return m.group(1)
    except Exception:
        pass
    return None


def get_all_postal_codes_on_page():
    """Get all postal codes visible on the current page"""
    try:
        return list(set(re.findall(r'\b(\d{4})\b', driver.page_source)))
    except Exception:
        return []


def clean_phone_number(phone_number):
    """Clean and format Norwegian phone number"""
    if not phone_number:
        return phone_number
    cleaned = re.sub(r'[^\d\+]', '', phone_number)
    if cleaned.startswith('0047'):
        cleaned = cleaned[4:]
    elif cleaned.startswith('+47'):
        cleaned = cleaned[3:]
    elif cleaned.startswith('47') and len(cleaned) == 10:
        cleaned = cleaned[2:]
    if len(cleaned) == 8 and cleaned.isdigit():
        return f"{cleaned[:2]} {cleaned[2:4]} {cleaned[4:6]} {cleaned[6:8]}"
    return cleaned


def _words(text):
    return set(re.findall(r"\w+", text.lower()))


def pick_matching_result(name, zip_code):
    """Return the /person/ link whose name matches `name` (words in any
    order). Returns None if nobody matches - never a wrong person."""
    name_list = re.findall(r"\w+", name.lower())
    if not name_list:
        return None
    wanted = set(name_list)
    first_last = {name_list[0], name_list[-1]}
    best = None  # (score, link)
    for a in driver.find_elements(By.CSS_SELECTOR, "a[href*='/person/']"):
        try:
            link_text = (a.text or a.get_attribute("textContent") or "").strip()
            words = _words(link_text)
            if wanted <= words:
                score = 2
            elif len(name_list) > 2 and first_last <= words:
                score = 1
            else:
                continue
            try:
                box = a.find_element(
                    By.XPATH, "./ancestor::*[self::li or self::article or self::div][1]"
                )
                if zip_code and zip_code in box.text:
                    score += 0.5
            except Exception:
                pass
            if best is None or score > best[0]:
                best = (score, a)
        except Exception:
            continue
    return best[1] if best else None


def log_page_links(limit=15):
    """Diagnostics: print what is actually on the page."""
    try:
        print(f"   URL: {driver.current_url}")
        print(f"   Title: {driver.title}")
        shown = 0
        for a in driver.find_elements(By.TAG_NAME, "a"):
            href = a.get_attribute("href") or ""
            text = (a.text or "").strip().replace("\n", " ")[:60]
            if href and text:
                print(f"   link: {text} -> {href[:100]}")
                shown += 1
                if shown >= limit:
                    break
    except Exception as e:
        print(f"   (could not list links: {e})")


def get_phone_number(name, zip_code, retry_count=0, max_retries=2):
    """Search 1881.no for a person and return the phone number (or '-')."""
    try:
        print(f"🔍 Searching for: {name} in postal code {zip_code}")

        search_input = WebDriverWait(driver, 10).until(
            EC.presence_of_element_located((By.CSS_SELECTOR, "input[type='search']"))
        )
        search_input.clear()
        search_input.send_keys(f"{name} {zip_code}")
        search_input.submit()
        time.sleep(3)

        current_url = driver.current_url
        print(f"📍 Current URL: {current_url}")

        # Search results page (1881 skips it when exactly one person matches)
        if "/person/" not in current_url:
            print("📋 On search results page, looking for a matching person...")
            try:
                try:
                    WebDriverWait(driver, 8).until(
                        EC.presence_of_element_located((By.CSS_SELECTOR, "a[href*='/person/']"))
                    )
                except TimeoutException:
                    pass

                match = pick_matching_result(name, zip_code)
                if match is None:
                    print("❌ No result matches the searched name")
                    log_page_links()
                    return "-"

                print(f"🔗 Matching result: {match.get_attribute('href')}")
                driver.execute_script("arguments[0].click();", match)
                time.sleep(3)

                current_url = driver.current_url
                print(f"📍 After click - Current URL: {current_url}")
                if "/person/" not in current_url:
                    print("❌ Failed to navigate to person detail page after clicking")
                    return "-"
            except Exception as e:
                print(f"❌ Error clicking matching result: {e}")
                return "-"

        if "/person/" in current_url:
            print("✅ Successfully navigated to person detail page!")

            # Verify the name on the page
            try:
                name_element = driver.find_element(By.CSS_SELECTOR, ".details-name, h1")
                page_name = name_element.text.strip()
                print(f"👤 Found person: {page_name}")
                if not verify_person_match(page_name, name, zip_code):
                    print("❌ Person verification failed")
                    return "-"
            except Exception:
                print("⚠️ Could not verify person name")

            # Verify the postal code
            try:
                address_elements = driver.find_elements(
                    By.CSS_SELECTOR, ".listing-address, .address, .location")
                address_found = False
                found_postal_codes = []
                for addr_elem in address_elements:
                    address_text = addr_elem.text
                    print(f"📍 Found address: {address_text}")
                    extracted_zip = extract_postal_code_from_address(address_text)
                    if extracted_zip:
                        found_postal_codes.append(extracted_zip)
                    if zip_code in address_text:
                        print(f"✅ Postal code {zip_code} confirmed in address!")
                        address_found = True
                        break
                if not address_found and zip_code in driver.page_source:
                    print(f"✅ Postal code {zip_code} found in page source!")
                    address_found = True
                if not address_found:
                    print(f"⚠️ Could not confirm postal code {zip_code}")
                    print(f"   Found postal codes on page: {found_postal_codes}")
                    return "-"
            except Exception:
                print("⚠️ Could not check address")

            print("📞 Looking for phone number...")
            page_postal_codes = get_all_postal_codes_on_page()
            print(f"📍 Postal codes found on page: {page_postal_codes}")

            # Strategy 1: tel: links
            try:
                phone_links = driver.find_elements(By.CSS_SELECTOR, "a[href^='tel:']")
                print(f"📞 Found {len(phone_links)} phone link(s)")
                for i, link in enumerate(phone_links):
                    phone_number = link.get_attribute('href').replace('tel:', '').strip()
                    if phone_number:
                        cleaned_phone = clean_phone_number(phone_number)
                        print(f"📞 Phone link {i+1}: {cleaned_phone}")
                        if zip_code and zip_code not in page_postal_codes:
                            print(f"⚠️ Skipping phone {cleaned_phone} - postal code {zip_code} not found on page")
                            continue
                        print(f"✅ Returning phone number: {cleaned_phone}")
                        return cleaned_phone
            except Exception as e:
                print(f"⚠️ Error processing phone links: {e}")

            # Strategy 2: specific elements
            phone_selectors = [
                ".button-call__number", ".contact-featured__call", ".phone",
                ".telephone", ".tlf", ".mobil", ".contact-phone",
            ]
            for selector in phone_selectors:
                try:
                    phone_elements = driver.find_elements(By.CSS_SELECTOR, selector)
                    print(f"📞 Found {len(phone_elements)} element(s) with selector '{selector}'")
                    for i, elem in enumerate(phone_elements):
                        phone_text = elem.text.strip()
                        if phone_text and re.search(r'\d', phone_text):
                            phone_clean = re.sub(r'[^\d\+\s]', '', phone_text)
                            if len(phone_clean.replace(' ', '')) >= 8:
                                cleaned_phone = clean_phone_number(phone_clean)
                                print(f"📞 Phone element {i+1}: {cleaned_phone}")
                                if zip_code and zip_code not in page_postal_codes:
                                    print(f"⚠️ Skipping phone {cleaned_phone} - postal code {zip_code} not found on page")
                                    continue
                                print(f"✅ Returning phone number: {cleaned_phone}")
                                return cleaned_phone
                except Exception as e:
                    print(f"⚠️ Error with selector '{selector}': {e}")
                    continue

            # Strategy 3: regex over the page source
            page_text = driver.page_source
            phone_patterns = [
                r'(\+47\s?)?(\d{2}\s?\d{2}\s?\d{2}\s?\d{2})',
                r'(\+47\s?)?\d{8}',
                r'tel:(\+47)?(\d{8})',
            ]
            for pattern in phone_patterns:
                for match in re.findall(pattern, page_text):
                    phone_number = ''.join(match) if isinstance(match, tuple) else match
                    phone_clean = re.sub(r'[^\d\+]', '', phone_number)
                    if len(phone_clean) >= 8:
                        cleaned_phone = clean_phone_number(phone_clean)
                        print(f"📞 Found phone number with regex: {cleaned_phone}")
                        return cleaned_phone

            print("❌ No phone number found on person page")
            return "-"

        print("❌ Did not reach person detail page")
        if retry_count < max_retries:
            print(f"🔄 Retrying with just name (attempt {retry_count + 1}/{max_retries})...")
            time.sleep(2)
            return get_phone_number(name, "", retry_count + 1, max_retries)
        return "-"

    except Exception as e:
        print(f"❌ Error during phone extraction: {e}")
        if retry_count < max_retries:
            print(f"🔄 Retrying after error (attempt {retry_count + 1}/{max_retries})...")
            time.sleep(3)
            return get_phone_number(name, zip_code, retry_count + 1, max_retries)
        return "-"


consent_handled = handle_consent_popup()
if not consent_handled:
    print("⚠️ Consent popup handling failed, but continuing...")

print("✅ Successfully loaded 1881.no and ready to extract phone numbers")

try:
    print(f"\n📊 Loading workbook: {file_path}")
    wb = load_workbook(file_path)
    print("✅ Workbook loaded successfully")
    print(f"📋 Available sheets: {wb.sheetnames}")
except Exception as e:
    print(f"❌ ERROR: Failed to load workbook: {e}")
    driver.quit()
    sys.exit(1)

for sheet_name in sheets:
    if sheet_name not in wb.sheetnames:
        print(f"⚠️ Sheet '{sheet_name}' not found in workbook, skipping...")
        continue

    print(f"\n📊 Processing sheet: {sheet_name}")
    print("=" * 50)

    try:
        df = pd.read_excel(file_path, sheet_name=sheet_name)
        print(f"📋 Found {len(df)} rows in sheet '{sheet_name}'")

        if "SHAREHOLDER" not in df.columns or "ZIP CODE / CITY" not in df.columns:
            print(f"⚠️ Required columns not found in sheet '{sheet_name}', skipping...")
            continue

        if "TELEPHONE" not in df.columns:
            df["TELEPHONE"] = ""
            print("📋 Added TELEPHONE column")

        phone_numbers = []

        for index, row in df.iterrows():
            name = str(row["SHAREHOLDER"]).strip()

            zip_city = str(row["ZIP CODE / CITY"]).strip()
            postcode_match = re.match(r'(\d{4})', zip_city)
            postcode = postcode_match.group(1) if postcode_match else zip_city[:4]
            print(f"📮 Extracted postal code: '{postcode}' from '{zip_city}'")

            if not name or name.lower() in ['nan', 'none', '']:
                phone_numbers.append("-")
                continue

            print(f"\n📋 Processing {index + 1}/{len(df)}: {name} (Zip: {postcode})")

            existing = str(row.get("TELEPHONE")).strip() if pd.notna(row.get("TELEPHONE")) else ""
            if existing and existing != "-":
                print(f"✅ Phone number already exists: {existing}")
                phone_numbers.append(existing)
                continue

            phone_number = get_phone_number(name, postcode)
            phone_numbers.append(phone_number)
            print(f"📞 Result: {phone_number if phone_number else 'No phone found'}")

            time.sleep(random.uniform(2, 4))

        ws = wb[sheet_name]
        telephone_col = None
        for col_idx, col in enumerate(df.columns, 1):
            if col == "TELEPHONE":
                telephone_col = col_idx
                break

        if telephone_col and len(phone_numbers) == len(df):
            for i, value in enumerate(phone_numbers, start=2):
                ws.cell(row=i, column=telephone_col, value=value)
            print(f"✅ Updated {len(phone_numbers)} phone numbers in sheet '{sheet_name}'")
        else:
            print(f"⚠️ Could not update sheet '{sheet_name}' - column mismatch")

    except Exception as e:
        print(f"❌ Error processing sheet '{sheet_name}': {e}")
        continue

try:
    wb.save(file_path)
    wb.close()
    print("\n💾 Workbook saved successfully")

    file_dir, file_name = os.path.split(file_path)
    file_base, file_ext = os.path.splitext(file_name)
    new_file_name = f"{file_base} (T){file_ext}"
    os.rename(file_path, os.path.join(file_dir, new_file_name))
    print(f"✅ SUCCESS: Phone numbers added and file saved as: {new_file_name}")
except Exception as e:
    print(f"❌ Error saving workbook: {e}")

driver.quit()
print("🔄 Browser closed. Process completed!")
